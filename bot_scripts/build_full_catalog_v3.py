"""
Build the FULL Aosom ES catalogue into Shopify import CSVs (v3 - variant-aware).

On top of v2 (brand scrub, copy repair, Shopify taxonomy mapping), this version
MERGES colour variants into a single parent product:

  * Group key  = SKU with the trailing colour code stripped (e.g. 831-135BK/WT/GY -> 831-135).
  * A group with >=2 distinct colours becomes ONE product:
        Option1 name = "Color", one variant row per colour (own SKU / price / cost / stock /
        weight / variant image).
  * Groups with a single colour (or a single member) stay a single product
        (Option1 "Title" / "Default Title"), keeping the colour word in the title.
  * Parent titles have the trailing colour word stripped so the colour lives on the variant.

Usage: python build_full_catalog_v3.py [xlsx] [out_dir] [--max-images 10]
"""
import os
import re
import csv
import json
import glob
import math
import html
import argparse
import collections
import warnings

from build_full_catalog_v2 import (
    scrub, pricing, slug, shopify_category, build_body, FIELDNAMES, SAFETY_BUFFER,
)

warnings.filterwarnings("ignore")

MAX_IMAGES = 10

# ---------------------------------------------------------------- colour handling

def normalize_color(c):
    c = (c or "").strip()
    return c or ""


def color_words(all_colors):
    ws = sorted({normalize_color(c).lower() for c in all_colors if c.strip()}, key=len, reverse=True)
    return [w for w in ws if len(w) >= 3]


def strip_color(title, words):
    """Drop a trailing colour phrase from a title."""
    low = title.lower()
    for w in words:
        if w and low.endswith(w) and len(title) > len(w) + 1:
            return title[:len(title) - len(w)].rstrip(" \t-–,;").strip()
    return title


def sku_base(sku):
    """Strip a trailing 2-4 uppercase-letter colour code from a SKU."""
    s = (sku or "").strip()
    if re.search(r"[A-Z]{2,4}$", s) and not re.search(r"[0-9]$", s):
        return re.sub(r"[A-Z]{2,4}$", "", s).rstrip("-")
    return s


# ---------------------------------------------------------------- build

def iter_catalog(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Product info"] if "Product info" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx = {h: i for i, h in enumerate(header) if h}
    for row in it:
        yield row, idx
    wb.close()


def row_title(row, idx):
    name = str(row[idx.get("Product Name (Spanish)")] or row[idx["Product Name"]] or "").strip()
    t = scrub(name, sentences=False, add_period=False)
    t = re.sub(r"^\s*nuestra colección[\s,:.-]*", "", t).strip()
    return re.sub(r"\s{2,}", " ", t)[:255]


def row_field(row, idx, name):
    if name in idx:
        return row[idx[name]]
    return ""


def money(v):
    if v is None:
        return 0.0
    s = str(v).replace("€", "").replace("\u00a0", "").replace(" ", "").strip()
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def row_images(row, idx, max_images):
    out = []
    for n in range(1, 13):
        u = str(row_field(row, idx, f"Product Image {n}") or "").strip()
        if u.startswith("http") and u not in out:
            out.append(u)
    return out[:max_images]


def make_variant(row, idx, colour, max_images):
    sku = str(row_field(row, idx, "Seller SKU") or "").strip()
    cost = money(row_field(row, idx, "List Price (Excl. Tax)"))
    qty = int(num(row_field(row, idx, "Qty Available")))
    grams = int(num(row_field(row, idx, "Product Weight (kg)")) * 1000)
    images = row_images(row, idx, max_images)
    return {
        "sku": re.sub(r"\s+", "-", sku).upper(),
        "colour": colour,
        "cost": cost,
        "price": pricing(cost)["price"],
        "compare": pricing(cost)["compare_at_price"],
        "avail": max(0, qty - SAFETY_BUFFER),
        "active": qty > SAFETY_BUFFER,
        "grams": grams,
        "images": images,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=sorted(glob.glob(
        os.path.expanduser("~/Downloads/Aosom ESStore Products *.xlsx")))[-1])
    ap.add_argument("out_dir", nargs="?", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..", "shopify_import_full"))
    ap.add_argument("--rows-per-file", type=int, default=20000)
    ap.add_argument("--max-images", type=int, default=10)
    args = ap.parse_args()
    os.makedirs(args.out_dir, exist_ok=True)
    # remove stale chunks from a previous build so the folder only holds this run's output
    for stale in glob.glob(os.path.join(args.out_dir, "shopify_import_es_full_*.csv")):
        os.remove(stale)
    for stale in ("manifest.json", "qa_report.json", "inventory_on_hand.csv", "preview.html",
                  "delivery_notes.txt"):
        sp = os.path.join(args.out_dir, stale)
        if os.path.exists(sp):
            os.remove(sp)

    # ---- collect rows
    catalog = []  # (sku, row, idx)
    all_colors = []
    for row, idx in iter_catalog(args.xlsx):
        sku = str(row_field(row, idx, "Seller SKU") or "").strip()
        if not sku:
            continue
        all_colors.append(row_field(row, idx, "Main Color"))
        catalog.append((sku, row, idx))

    cwords = color_words(all_colors)

    # ---- group by SKU base, preserving order
    groups = collections.OrderedDict()
    for sku, row, idx in catalog:
        groups.setdefault(sku_base(sku), []).append((sku, row, idx))

    products = []          # list of parent dicts for QA
    rows_out = []          # csv dict rows
    seen_skus = set()
    costs, prices = [], []

    def emit_product(rec, variants, shared_images, max_images):
        """rec = parent-level dict; variants = list of make_variant dicts."""
        handle = rec["handle"]
        any_active = any(v["active"] for v in variants)
        for i, v in enumerate(variants):
            if v["sku"] in seen_skus:
                continue
            seen_skus.add(v["sku"])
            r = {
                "Title": rec["title"],
                "URL handle": handle,
                "Description": rec["body"] if i == 0 else "",
                "Vendor": rec["vendor"],
                "Product category": rec["category"] if i == 0 else "",
                "Type": rec["type"] if i == 0 else "",
                "Published on online store": "TRUE" if i == 0 else "",
                "Status": "active" if i == 0 else "",
                "SKU": v["sku"],
                "Option1 name": rec["opt_name"],
                "Option1 value": v["colour"] or "Default Title",
                "Price": f"{v['price']:.2f}",
                "Compare-at price": "",
                "Cost per item": f"{v['cost']:.2f}",
                "Charge tax": "TRUE",
                "Inventory tracker": "shopify",
                "Inventory quantity": str(v["avail"]),
                "Continue selling when out of stock": "deny",
                "Weight value (grams)": str(v["grams"]),
                "Weight unit for display": "kg",
                "Requires shipping": "TRUE",
                "Fulfillment service": "manual",
                "Product image URL": (shared_images[0] if shared_images else "") if i == 0 else "",
                "Image position": "1" if i == 0 else "",
                "Image alt text": rec["title"] if i == 0 else "",
                "Variant image URL": v["images"][0] if v["images"] else "",
                "Gift card": "FALSE" if i == 0 else "",
                "SEO title": rec["seo_title"] if i == 0 else "",
                "SEO description": rec["seo_desc"] if i == 0 else "",
                "Google Shopping / Manufacturer part number (MPN)": v["sku"],
                "Google Shopping / Condition": "new" if i == 0 else "",
                "Google Shopping / Custom product": "FALSE" if i == 0 else "",
                "Google Shopping / Custom label 0": "",
            }
            rows_out.append(r)
            costs.append(v["cost"])
            prices.append(v["price"])
        # product gallery (positions 2..N) — position 1 is the first variant's Product image URL
        for n, url in enumerate(shared_images[1:], start=2):
            rows_out.append({"URL handle": handle, "Product image URL": url,
                             "Image position": str(n),
                             "Image alt text": f"{rec['title']} - vista {n}"})

    for base, members in groups.items():
        # normalise colour per member
        cols = [normalize_color(row_field(row, idx, "Main Color")) for _, row, idx in members]
        distinct = {c for c in cols if c}
        is_variant_group = len(members) > 1 and len(distinct) >= 2

        if is_variant_group:
            variants = []
            used_colours = {}
            for (sku, row, idx), c in zip(members, cols):
                label = c or "Color"
                if label in used_colours:
                    label = f"{c} ({re.sub(r'[^A-Z0-9]', '', sku[-3:]).upper()})"
                used_colours[label] = True
                variants.append(make_variant(row, idx, label, args.max_images))
            first = members[0]
            first_title = row_title(first[1], first[2])
            parent_title = strip_color(first_title, cwords) or first_title
            cat_en = str(row_field(first[1], first[2], "Category") or "").strip()
            body = build_body(first[1], first[2])
            shared = row_images(first[1], first[2], args.max_images)
            handle = f"{slug(base)}-{slug(parent_title)}"
            rec = {
                "handle": handle, "title": parent_title, "body": body, "vendor": "",
                "category": shopify_category(cat_en), "type": cat_en,
                "tags": ", ".join(dict.fromkeys(["Envío desde Europa", "Stock verificado", cat_en])),
                "opt_name": "Color",
                "seo_title": parent_title[:70],
                "seo_desc": (parent_title + " - Envío en 3-5 días desde almacén europeo. Garantía 12 meses.")[:160],
            }
            emit_product(rec, variants, shared, args.max_images)
            products.append({"handle": handle, "title": parent_title,
                             "variants": len(variants), "type": cat_en, "base": base})
        else:
            for sku, row, idx in members:
                title = row_title(row, idx)
                if not title:
                    continue
                v = make_variant(row, idx, "Default Title", args.max_images)
                cat_en = str(row_field(row, idx, "Category") or "").strip()
                body = build_body(row, idx)
                shared = row_images(row, idx, args.max_images)
                handle = f"{slug(v['sku'])}-{slug(title)}"
                rec = {
                    "handle": handle, "title": title, "body": body, "vendor": "",
                    "category": shopify_category(cat_en), "type": cat_en,
                    "tags": ", ".join(dict.fromkeys(["Envío desde Europa", "Stock verificado", cat_en])),
                    "opt_name": "Title",
                    "seo_title": title[:70],
                    "seo_desc": (title + " - Envío en 3-5 días desde almacén europeo. Garantía 12 meses.")[:160],
                }
                emit_product(rec, [v], shared, args.max_images)
                products.append({"handle": handle, "title": title, "variants": 1, "type": cat_en,
                                 "base": v["sku"]})

    # ---- de-duplicate parent titles (append SKU base when they collide)
    title_count = collections.Counter(p["title"] for p in products)
    dup_titles = {t for t, c in title_count.items() if c > 1}
    fixed_titles = 0
    for p in products:
        if p["title"] in dup_titles:
            p["title"] = f"{p['title']} - {slug(p['base']).upper()}"[:255]
            fixed_titles += 1
    # apply the (possibly changed) titles back into csv rows
    by_handle = {p["handle"]: p for p in products}
    for r in rows_out:
        h = r.get("URL handle")
        if h and r.get("Title") and h in by_handle:
            t = by_handle[h]["title"]
            r["Title"] = t
            r["Image alt text"] = t
            r["SEO title"] = t[:70]
            r["SEO description"] = (t + " - Envío en 3-5 días desde almacén europeo. Garantía 12 meses.")[:160]

    # ---- write chunks
    files = []
    n_chunks = math.ceil(len(rows_out) / args.rows_per_file)
    total = 0
    for c in range(n_chunks):
        chunk = rows_out[c * args.rows_per_file:(c + 1) * args.rows_per_file]
        path = os.path.join(args.out_dir, f"shopify_import_es_full_{c + 1:02d}.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=FIELDNAMES)
            w.writeheader()
            for r in chunk:
                w.writerow({k: r.get(k, "") for k in FIELDNAMES})
        files.append({"file": os.path.basename(path), "rows": len(chunk) + 1, "bytes": os.path.getsize(path)})
        total += len(chunk) + 1

    # ---- QA
    leaks = []
    for r in rows_out:
        if not r.get("Title"):
            continue
        blob = " ".join(str(r.get(k, "")) for k in
                        ("Title", "Description", "Tags", "SEO title", "SEO description")).lower()
        hit = next((b for b in ["aosom", "homcom", "outsunny", "pawhut", "vinsetto", "sobuy",
                                "songmics", "arkswift", "alibaba", "aliexpress"] if b in blob), None)
        if hit:
            leaks.append({"sku": r.get("SKU"), "leak": hit})

    variant_parents = sum(1 for p in products if p["variants"] > 1)
    total_variants = sum(p["variants"] for p in products)
    manifest = {
        "source": args.xlsx,
        "parents": len(products),
        "variant_parents": variant_parents,
        "total_variants": total_variants,
        "single_products": len(products) - variant_parents,
        "csv_rows": total,
        "files": files,
        "total_bytes": sum(f["bytes"] for f in files),
        "duplicate_titles_fixed": fixed_titles,
        "whitewash_leaks": leaks,
        "price_min": min(prices) if prices else 0,
        "price_max": max(prices) if prices else 0,
        "avg_margin_pct": round(sum((s - c) / s * 100 for c, s in zip(costs, prices)) / len(prices), 1) if prices else 0,
        "template": "shopify-current (variants via Option1 name=Color)",
    }
    json.dump(manifest, open(os.path.join(args.out_dir, "manifest.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    print(json.dumps({k: manifest[k] for k in
                      ("parents", "variant_parents", "total_variants", "single_products",
                       "csv_rows", "total_bytes", "duplicate_titles_fixed",
                       "whitewash_leaks", "price_min", "price_max", "avg_margin_pct", "files")},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
