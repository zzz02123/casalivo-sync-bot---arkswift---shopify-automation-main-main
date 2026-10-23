"""
Build the FULL Aosom ES catalogue into Shopify import CSVs — NO white-label brand,
matching Shopify's CURRENT product CSV template (the new header names, e.g.
"URL handle", "Description", "SKU", "Inventory quantity").

Source: the complete Aosom ESStore Products export (8156 rows) — self-contained.
Usage: python build_full_catalog.py [xlsx] [out_dir] [--max-per-file 2000]
"""
import os
import re
import csv
import json
import glob
import math
import html
import unicodedata
import argparse
import warnings
from typing import Any, Dict, List, Optional

warnings.filterwarnings("ignore")

SAFETY_BUFFER = 3

BANNED = [
    "aosom", "homcom", "outsunny", "pawhut", "vinsetto", "sobuy", "songmics",
    "songesang", "happybird", "spain aosom", "spanish aosom sl",
    "tienda oficial", "distribuidor oficial", "garantía del fabricante",
    "manufacturer warranty", "official store", "spanish warehouse", "almacén españa",
    "aosom.es", "arkswift", "alibaba", "aliexpress",
]

CATEGORY_ES = {
    "greenhouse": "Invernaderos y Jardín", "garden & outdoor": "Jardín y Exterior",
    "garden": "Jardín y Exterior", "outdoor": "Jardín y Exterior",
    "furniture": "Muebles", "home furniture": "Muebles", "office": "Oficina y Despacho",
    "home & living": "Hogar", "home": "Hogar", "storage": "Almacenaje y Organización",
    "pet supplies": "Mascotas", "pets": "Mascotas", "sports & outdoor": "Deporte y Aire Libre",
    "toys": "Juguetes", "kitchen": "Cocina", "bathroom": "Baño", "bedroom": "Dormitorio",
    "lighting": "Iluminación", "tools": "Herramientas", "appliance": "Electrodomésticos",
}

GROUP_RULES = [
    ("mascota", ("cat ", "dog ", "pet", "litter", "scratcher")),
    ("oficina", ("office", "desk", "chair")),
    ("jardín", ("garden", "greenhouse", "shed", "grill", "parasol", "outdoor", "patio", "bbq", "planter", "gazebo", "awning", "marquee", "camping")),
    ("almacenaje", ("storage", "cabinet", "shelf", "shelves", "bookcase", "trolley", "jewellery", "wardrobe", "rack", "organizer")),
    ("muebles", ("table", "sofa", "bed", "nightstand", "dressing", "bench", "stool", "buffet", "sideboard", "chaise", "chair", "armchair", "furniture")),
]

# Shopify's current product CSV template header order (subset we populate)
FIELDNAMES = [
    "Title", "URL handle", "Description", "Vendor", "Product category", "Type", "Tags",
    "Published on online store", "Status", "SKU", "Option1 name", "Option1 value",
    "Price", "Compare-at price", "Cost per item", "Charge tax", "Inventory tracker",
    "Inventory quantity", "Continue selling when out of stock", "Weight value (grams)",
    "Weight unit for display", "Requires shipping", "Fulfillment service",
    "Product image URL", "Image position", "Image alt text", "Variant image URL",
    "Gift card", "SEO title", "SEO description",
    "Google Shopping / Manufacturer part number (MPN)", "Google Shopping / Condition",
    "Google Shopping / Custom product", "Google Shopping / Custom label 0",
]


def money(value: Any) -> float:
    if value is None:
        return 0.0
    s = str(value).replace("€", "").replace("\u00a0", "").replace(" ", "").strip()
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def num(value: Any) -> float:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def scrub(text: Any) -> str:
    if not text:
        return ""
    out = str(text)
    out = re.sub(r"https?://\S+", "", out)
    out = re.sub(r"[\w.+-]+@[\w-]+\.[\w.]+", "", out)
    for bad in BANNED:
        out = re.sub(rf"\b{re.escape(bad)}\b", "", out, flags=re.IGNORECASE)
    out = re.sub(r"[ \t]{2,}", " ", out)
    out = re.sub(r"\s+([,.;:])", r"\1", out)
    out = re.sub(r"^(?:\s*[-–|,;:.])+|(?:\s*[-–|,;:.])+$", "", out)
    return out.strip(" -–|,;\t\n")


def pricing(cost: float) -> Dict[str, float]:
    if cost <= 0:
        return {"price": 0.0, "compare_at_price": 0.0}
    mult = 2.5 if cost < 20 else (2.0 if cost <= 50 else 1.8)
    raw = cost * mult
    price = math.floor(raw) + 0.99
    if price < raw:
        price += 1.0
    compare = math.floor(price * 1.35) + 0.99
    return {"price": round(price, 2), "compare_at_price": round(compare, 2)}


def slug(text: str, limit: int = 55) -> str:
    norm = unicodedata.normalize("NFKD", text.lower()).encode("ASCII", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", norm).strip("-")[:limit].strip("-")


def spanish_group(category_en: str) -> str:
    low = (category_en or "").lower()
    for group, keys in GROUP_RULES:
        if any(k in low for k in keys):
            return group
    return CATEGORY_ES.get(low, "")


def feature_list(raw: Any) -> List[str]:
    if not raw:
        return []
    parts = [p.strip(" -•\t") for p in re.split(r"[\r\n]+", str(raw)) if p.strip()]
    return [scrub(p) for p in parts if len(scrub(p)) > 3][:8]


def spec_rows(row: List[Any], idx: Dict[str, int]) -> List[tuple]:
    specs: List[tuple] = []
    for label, col in [("Color", "Main Color"), ("Material", "Main Material"), ("Estilo", "Product Style")]:
        val = scrub(row[idx.get(col, -1)] or "") if col in idx else ""
        if val and val.lower() != "none":
            specs.append((label, val))
    dims = []
    for label, col in [("L", "Assembled Length (cm)"), ("An", "Assembled Width (cm)"), ("Al", "Assembled Height (cm)")]:
        v = num(row[idx.get(col, -1)]) if col in idx else 0
        if v > 0:
            dims.append(f"{label} {v:g}")
    if dims:
        specs.append(("Medidas (cm)", " x ".join(dims)))
    weight = num(row[idx.get("Product Weight (kg)", -1)]) if "Product Weight (kg)" in idx else 0
    if weight > 0:
        specs.append(("Peso", f"{weight:g} kg"))
    return specs


def build_body(row: List[Any], idx: Dict[str, int]) -> str:
    desc = scrub(row[idx.get("Description", -1)] or "")
    feats = feature_list(row[idx.get("Features", -1)] if "Features" in idx else "")
    parts: List[str] = []
    if desc:
        parts.append(f"<p>{html.escape(desc)}</p>")
    if feats:
        items = []
        for f in feats:
            if ":" in f:
                head, tail = f.split(":", 1)
                if len(head) < 60:
                    items.append(f"<li><strong>{html.escape(head.strip())}:</strong>{html.escape(tail)}</li>")
                    continue
            items.append(f"<li>{html.escape(f)}</li>")
        parts.append("<h3>Características principales</h3><ul>" + "".join(items) + "</ul>")
    specs = spec_rows(row, idx)
    if specs:
        rows_html = "".join(
            f"<tr><td><strong>{html.escape(k)}</strong></td><td>{html.escape(v)}</td></tr>" for k, v in specs
        )
        parts.append('<h3>Ficha técnica</h3><table style="width:100%;border-collapse:collapse">'
                     f"<tbody>{rows_html}</tbody></table>")
    parts.append(
        '<div style="margin-top:20px;padding:15px;background:#f1f5f9;border-radius:6px;">'
        "<p><strong>Envío y servicio</strong></p>"
        "<ul><li>Envío desde almacén europeo: 3-5 días laborables.</li>"
        "<li>Stock verificado en tiempo real.</li>"
        "<li>Seguimiento de pedido incluido.</li></ul></div>"
    )
    parts.append(
        '<div style="margin-top:16px;padding:15px;background:#f8fafc;border-left:4px solid #0f172a;">'
        "<p><strong>Garantía:</strong> 30 días para devoluciones y 12 meses de garantía "
        "en defectos de fabricación. Atención al cliente en español.</p></div>"
    )
    return "\n".join(parts)


def iter_catalog(xlsx_path: str):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Product info"] if "Product info" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx = {h: i for i, h in enumerate(header) if h}
    for row in it:
        yield row, idx


def build_row(row: List[Any], idx: Dict[str, int], max_images: int) -> Optional[Dict[str, str]]:
    sku = str(row[idx["Seller SKU"]] or "").strip()
    name = str(row[idx.get("Product Name (Spanish)")] or row[idx["Product Name"]] or "").strip()
    if not sku or not name:
        return None
    cost = money(row[idx["List Price (Excl. Tax)"]])
    qty = int(num(row[idx.get("Qty Available")]))
    images: List[str] = []
    for n in range(1, 13):
        url = str(row[idx.get(f"Product Image {n}")] or "").strip()
        if url.startswith("http") and url not in images:
            images.append(url)
    if not images:
        return None
    images = images[:max_images]

    title = scrub(name)
    if not title:
        return None
    sku_clean = re.sub(r"\s+", "-", sku).upper()
    handle = f"{slug(sku_clean)}-{slug(name)}"
    price = pricing(cost)
    avail = max(0, qty - SAFETY_BUFFER)
    is_active = qty > SAFETY_BUFFER
    grams = int(num(row[idx.get("Product Weight (kg)")]) * 1000)
    category_en = str(row[idx.get("Category")] or "").strip()
    category_es = spanish_group(category_en)
    body = build_body(row, idx)
    tags = [t for t in ["Envío desde Europa", "Stock verificado", category_es or category_en] if t]
    seo_title = title[:70]
    seo_desc = (scrub(name) + " - Envío en 3-5 días desde almacén europeo. Garantía 12 meses.")[:160]

    return {
        "Title": title,
        "URL handle": handle,
        "Description": body,
        "Vendor": "",
        "Product category": "",
        "Type": category_en,
        "Tags": ", ".join(dict.fromkeys(tags)),
        "Published on online store": "TRUE" if is_active else "FALSE",
        "Status": "active" if is_active else "draft",
        "SKU": sku_clean,
        "Option1 name": "Title",
        "Option1 value": "Default Title",
        "Price": f"{price['price']:.2f}",
        "Compare-at price": f"{price['compare_at_price']:.2f}",
        "Cost per item": f"{cost:.2f}",
        "Charge tax": "TRUE",
        "Inventory tracker": "shopify",
        "Inventory quantity": str(avail),
        "Continue selling when out of stock": "deny",
        "Weight value (grams)": str(grams),
        "Weight unit for display": "kg",
        "Requires shipping": "TRUE",
        "Fulfillment service": "manual",
        "Product image URL": images[0],
        "Image position": "1",
        "Image alt text": title,
        "Variant image URL": images[0],
        "Gift card": "FALSE",
        "SEO title": seo_title,
        "SEO description": seo_desc,
        "Google Shopping / Manufacturer part number (MPN)": sku_clean,
        "Google Shopping / Condition": "new",
        "Google Shopping / Custom product": "FALSE",
        "Google Shopping / Custom label 0": "",
        "_images": images[1:],
        "_title": title,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=sorted(glob.glob(
        os.path.expanduser("~/Downloads/Aosom ESStore Products *.xlsx")))[-1])
    ap.add_argument("out_dir", nargs="?", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..", "shopify_import_full"))
    ap.add_argument("--max-per-file", type=int, default=2000)
    ap.add_argument("--max-images", type=int, default=10)
    ap.add_argument("--min-qty", type=int, default=0)
    args = ap.parse_args()

    os.makedirs(args.out_dir, exist_ok=True)

    products: List[Dict[str, str]] = []
    seen_skus = set()
    costs, prices = [], []

    for row, idx in iter_catalog(args.xlsx):
        sku = str(row[idx["Seller SKU"]] or "").strip()
        if not sku or sku in seen_skus:
            continue
        if int(num(row[idx.get("Qty Available")])) < args.min_qty:
            continue
        main = build_row(row, idx, args.max_images)
        if not main:
            continue
        seen_skus.add(sku)
        products.append(main)
        costs.append(float(main["Cost per item"]))
        prices.append(float(main["Price"]))

    files = []
    n_chunks = math.ceil(len(products) / args.max_per_file)
    total_csv_rows = 0
    for c in range(n_chunks):
        chunk = products[c * args.max_per_file:(c + 1) * args.max_per_file]
        path = os.path.join(args.out_dir, f"shopify_import_es_full_{c + 1:02d}.csv")
        rows_written = 0
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=FIELDNAMES)
            w.writeheader()
            rows_written += 1
            for p in chunk:
                w.writerow({k: v for k, v in p.items() if not k.startswith("_")})
                rows_written += 1
                for n, url in enumerate(p["_images"], start=2):
                    w.writerow({"URL handle": p["URL handle"], "Product image URL": url,
                                "Image position": str(n),
                                "Image alt text": f"{p['_title']} - vista {n}"})
                    rows_written += 1
        files.append({"file": os.path.basename(path), "products": len(chunk),
                      "rows": rows_written, "bytes": os.path.getsize(path)})
        total_csv_rows += rows_written

    leaks = []
    for p in products:
        blob = " ".join(str(p.get(k, "")) for k in ("Title", "Description", "Tags", "SEO title", "SEO description")).lower()
        hit = next((b for b in BANNED + ["casalivo"] if b in blob), None)
        if hit:
            leaks.append({"sku": p.get("SKU"), "leak": hit})

    manifest = {
        "source": args.xlsx,
        "products": len(products),
        "active": sum(1 for p in products if p["Status"] == "active"),
        "draft": sum(1 for p in products if p["Status"] == "draft"),
        "csv_rows": total_csv_rows,
        "files": files,
        "total_bytes": sum(f["bytes"] for f in files),
        "whitewash_leaks": leaks,
        "price_min": min(prices) if prices else 0,
        "price_max": max(prices) if prices else 0,
        "avg_margin_pct": round(sum((s - c) / s * 100 for c, s in zip(costs, prices)) / len(prices), 1) if prices else 0,
        "brand": None,
        "safety_buffer": SAFETY_BUFFER,
        "template": "shopify-current (URL handle / Description / SKU / Inventory quantity)",
    }
    json.dump(manifest, open(os.path.join(args.out_dir, "manifest.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    print(json.dumps({k: manifest[k] for k in
                      ("products", "active", "draft", "csv_rows", "total_bytes",
                       "price_min", "price_max", "avg_margin_pct", "whitewash_leaks", "files")},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
