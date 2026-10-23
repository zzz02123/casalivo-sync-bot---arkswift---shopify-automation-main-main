"""
Casalivo - Real ArkSwift ES inventory -> Shopify product import builder
--------------------------------------------------------------------
Inputs (real ArkSwift export files):
  1) Inventory feed  : "库存 - 工作表1.csv"   (Store Name=Aosom ES, List Price "61,88€", Qty Available)
  2) Product catalog : "Aosom ESStore Products *.xlsx" (Spanish title, description, features, 12 images)

Outputs:
  * shopify_import_casalivo_es.csv   -> Shopify Products import (Admin UI > Products > Import)
  * shopify_import_casalivo_es.json  -> ready-to-POST Admin API payloads
  * import_qa_report.json            -> QA: counts, price/margin stats, whitewash scan

Usage:
  python build_casalivo_import.py [inventory.csv] [catalog.xlsx] [out_dir]
"""
import os
import re
import csv
import sys
import json
import glob
import html
import math
import unicodedata
import warnings
from typing import Any, Dict, List, Optional

warnings.filterwarnings("ignore")

BRAND = "Casalivo"
SKU_PREFIX = "CSL-"
SAFETY_BUFFER = 3
MAX_IMAGES = 10

# ---------------------------------------------------------------- whitewash
REPLACEMENTS = [
    (r"SPANISH AOSOM SL", BRAND), (r"AOSOM\.ES", "casalivo.es"), (r"AOSOM", BRAND),
    (r"HOMCOM", BRAND), (r"OUTSUNNY", BRAND), (r"PAWHUT", "Casalivo Pets"),
    (r"VINSETTO", "Casalivo Office"), (r"SOBUY", BRAND), (r"SONGMICS", BRAND),
]
BANNED = [
    "aosom", "homcom", "outsunny", "pawhut", "vinsetto", "sobuy", "songmics",
    "tienda oficial", "distribuidor oficial", "garantía del fabricante",
    "manufacturer warranty", "official store", "spanish warehouse", "almacén españa",
    "arkswift", "alibaba", "aliexpress",
]

CATEGORY_ES = {
    "greenhouse": "Invernaderos y Jardín",
    "garden & outdoor": "Jardín y Exterior",
    "garden": "Jardín y Exterior",
    "outdoor": "Jardín y Exterior",
    "furniture": "Muebles",
    "home furniture": "Muebles",
    "office": "Oficina y Despacho",
    "home & living": "Hogar",
    "home": "Hogar",
    "storage": "Almacenaje y Organización",
    "pet supplies": "Mascotas",
    "pets": "Mascotas",
    "sports & outdoor": "Deporte y Aire Libre",
    "toys": "Juguetes",
    "kitchen": "Cocina",
    "bathroom": "Baño",
    "bedroom": "Dormitorio",
    "lighting": "Iluminación",
    "tools": "Herramientas",
    "appliance": "Electrodomésticos",
}



GROUP_RULES = [
    ("mascota", ("cat ", "dog ", "pet", "litter", "scratcher")),
    ("oficina", ("office", "desk", "chair")),
    ("jardín", ("garden", "greenhouse", "shed", "grill", "parasol", "outdoor", "patio", "bbq", "planter")),
    ("almacenaje", ("storage", "cabinet", "shelf", "shelves", "bookcase", "trolley", "jewellery", "wardrobe", "rack")),
    ("muebles", ("table", "sofa", "bed", "nightstand", "dressing", "bench", "stool", "buffet", "sideboard", "chaise")),
]


SHOPIFY_TAXONOMY = {
    "greenhouse": "Home & Garden > Lawn & Garden > Gardening > Greenhouses",
    "nightstands": "Furniture > Bedroom Furniture > Nightstands",
    "storage trolleys": "Furniture > Carts & Islands > Kitchen & Dining Carts",
    "trash cans": "Home & Garden > Household Supplies > Waste Containment > Trash Cans",
    "sofa beds & chaise lounges": "Furniture > Living Room Furniture > Sofas",
    "storage cabinets": "Furniture > Cabinets & Storage > Storage Cabinets & Lockers",
    "awnings": "Hardware > Building Materials > Canopies & Awnings",
    "garden furniture sets": "Furniture > Outdoor Furniture > Outdoor Furniture Sets",
    "sun loungers": "Furniture > Outdoor Furniture > Outdoor Chairs > Sunloungers",
    "parasol umbrellas & accessories": "Home & Garden > Lawn & Garden > Outdoor Living > Outdoor Umbrellas & Sunshades",
    "tables": "Furniture > Tables > Kitchen & Dining Room Tables",
    "shade sails": "Home & Garden > Lawn & Garden > Outdoor Living > Outdoor Umbrellas & Sunshades",
    "charcoal grills": "Home & Garden > Lawn & Garden > Outdoor Living > Outdoor Grills",
    "bookcases & shelves": "Furniture > Shelving > Bookcases & Standing Shelves",
    "camping & hiking": "Sporting Goods > Outdoor Recreation > Camping & Hiking",
    "pet supplies": "Animals & Pet Supplies > Pet Supplies",
}


def shopify_standard_category(category_en: str) -> str:
    low = (category_en or "").lower().strip()
    if low in SHOPIFY_TAXONOMY:
        return SHOPIFY_TAXONOMY[low]
    for k, v in SHOPIFY_TAXONOMY.items():
        if k in low or low in k:
            return v
    if any(w in low for w in ("garden", "greenhouse", "outdoor", "awning", "parasol")):
        return "Home & Garden > Lawn & Garden"
    if any(w in low for w in ("furniture", "chair", "table", "bed", "sofa", "shelf", "cabinet")):
        return "Furniture > Living Room Furniture"
    if any(w in low for w in ("cat", "dog", "pet")):
        return "Animals & Pet Supplies > Pet Supplies"
    return "Home & Garden > Lawn & Garden"


def spanish_group(category_en: str) -> str:
    low = (category_en or "").lower()
    for group, keys in GROUP_RULES:
        if any(k in low for k in keys):
            return group
    return CATEGORY_ES.get(low, "Hogar")


def money(value: str) -> float:
    """'1.234,56€' / '61,88€' -> float"""
    if value is None:
        return 0.0
    s = str(value).replace("€", "").replace("\u00a0", "").replace(" ", "").strip()
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def num(value: Any) -> float:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def scrub(text: str) -> str:
    """Remove every supplier trace from a string."""
    if not text:
        return ""
    out = str(text)
    for pat, rep in REPLACEMENTS:
        out = re.sub(pat, rep, out, flags=re.IGNORECASE)
    out = re.sub(r"https?://\S+", "", out)
    out = re.sub(r"[\w.+-]+@[\w-]+\.[\w.]+", "", out)
    for bad in BANNED:
        out = re.sub(rf"\b{re.escape(bad)}\b", "", out, flags=re.IGNORECASE)
    out = re.sub(r"[ \t]{2,}", " ", out)
    out = re.sub(r"\s+([,.;:])", r"\1", out)
    return out.strip(" -–|,;\t\n")


def slug(text: str, limit: int = 55) -> str:
    norm = unicodedata.normalize("NFKD", text.lower()).encode("ASCII", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", norm).strip("-")[:limit].strip("-")


def pricing(cost: float) -> Dict[str, float]:
    if cost <= 0:
        return {"price": 0.0, "compare_at_price": 0.0, "profit": 0.0, "margin": 0.0}
    if cost <= 100:
        mult = 1.55
    elif cost <= 200:
        mult = 1.50
    else:
        mult = 1.45
    raw = cost * mult
    price = round(raw, 2)
    compare = round(price * 1.35, 2)
    return {
        "price": round(price, 2),
        "compare_at_price": round(compare, 2),
        "profit": round(price - cost, 2),
        "margin": round((price - cost) / price * 100, 1) if price else 0.0,
    }


def generate_ean13_barcode(seed: str = None) -> str:
    """生成13位EAN码，前缀 ean:，国家码 841，厂商 57926，随机 4位 + Modulo-10 校验"""
    import hashlib, random
    if seed:
        h = int(hashlib.md5(str(seed).encode("utf-8")).hexdigest()[:8], 16)
        num = h % 10000
    else:
        num = random.randint(0, 9999)
    item_code = f"{num:04d}"
    raw12 = f"84157926{item_code}"
    s = sum(int(ch) * (1 if i % 2 == 0 else 3) for i, ch in enumerate(raw12))
    check = (10 - (s % 10)) % 10
    return f"ean:{raw12}{check}"


def clean_title(raw: str) -> str:
    title = scrub(raw)
    title = re.sub(r"\s+", " ", title).strip()
    if not title:
        return ""
    if not title.lower().startswith(BRAND.lower()):
        title = f"{BRAND} {title}"
    return title[:255]


def feature_list(raw: str) -> List[str]:
    if not raw:
        return []
    parts = [p.strip(" -•\t") for p in re.split(r"[\r\n]+", str(raw)) if p.strip()]
    return [scrub(p) for p in parts if len(scrub(p)) > 3][:8]


def spec_rows(row: Dict[str, Any], idx: Dict[str, int]) -> List[tuple]:
    specs = []
    for label, col in [
        ("Color", "Main Color"),
        ("Material", "Main Material"),
        ("Estilo", "Product Style"),
    ]:
        val = scrub(str(row[idx.get(col, -1)] or "")) if col in idx else ""
        if val and val.lower() != "none":
            specs.append((label, val))
    dims = []
    for label, col in [("L", "Assembled Length (cm)"), ("An", "Assembled Width (cm)"), ("Al", "Assembled Height (cm)")]:
        v = num(row[idx.get(col, -1)]) if col in idx else 0
        if v > 0:
            dims.append(f"{label} {v:g}")
    if dims:
        specs.append(("Medidas (cm)", " x ".join(dims).replace("L ", "").replace("An ", "").replace("Al ", "")))
    weight = num(row[idx.get("Product Weight (kg)", -1)]) if "Product Weight (kg)" in idx else 0
    if weight > 0:
        specs.append(("Peso", f"{weight:g} kg"))
    return specs


def build_body(row: Dict[str, Any], idx: Dict[str, int], sku: str) -> str:
    desc = scrub(str(row[idx.get("Description", -1)] or ""))
    feats = feature_list(row[idx.get("Features", -1)] if "Features" in idx else "")
    parts = []
    if desc:
        parts.append(f"<p>{html.escape(desc)}</p>")
    if feats:
        items = []
        for f in feats:
            if ":" in f:
                head, tail = f.split(":", 1)
                if len(head) < 60:
                    items.append(f"<li><strong>{html.escape(head.strip())}:</strong>{html.escape(tail)}</li>")
                    continue
            items.append(f"<li>{html.escape(f)}</li>")
        parts.append("<h3>Características principales</h3><ul>" + "".join(items) + "</ul>")
    specs = spec_rows(row, idx)
    if specs:
        rows_html = "".join(f"<tr><td><strong>{html.escape(k)}</strong></td><td>{html.escape(v)}</td></tr>" for k, v in specs)
        parts.append(
            '<h3>Ficha técnica</h3><table style="width:100%;border-collapse:collapse">'
            f'<tbody>{rows_html}</tbody></table>'
        )
    body = "\n".join(parts)
    return body


def load_inventory(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def load_catalog(path: str) -> Dict[str, Dict[str, int]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Product info"] if "Product info" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx = {h: i for i, h in enumerate(header) if h}
    out = {}
    for row in it:
        sku = str(row[idx["Seller SKU"]] or "").strip()
        if sku and sku not in out:
            out[sku] = {"row": row, "idx": idx}
    return out


def main() -> None:
    inv_path = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Downloads/库存 - 工作表1.csv")
    cat_path = sys.argv[2] if len(sys.argv) > 2 else sorted(
        glob.glob(os.path.expanduser("~/Downloads/Aosom ESStore Products *.xlsx"))
    )[-1]
    out_dir = sys.argv[3] if len(sys.argv) > 3 else os.path.dirname(os.path.abspath(__file__)) + "/.."

    inventory = load_inventory(inv_path)
    catalog = load_catalog(cat_path)

    fieldnames = [
        "Handle", "Title", "Body (HTML)", "Vendor", "Product Category", "Type", "Tags",
        "Published", "Option1 Name", "Option1 Value", "Option2 Name", "Option2 Value",
        "Option3 Name", "Option3 Value", "Variant SKU", "Variant Grams",
        "Variant Inventory Tracker", "Variant Inventory Qty", "Variant Inventory Policy",
        "Variant Fulfillment Service", "Variant Price", "Variant Compare At Price",
        "Variant Requires Shipping", "Variant Taxable", "Variant Barcode",
        "Image Src", "Image Position", "Image Alt Text", "Gift Card",
        "SEO Title", "SEO Description", "Google Shopping / MPN",
        "Google Shopping / Condition", "Google Shopping / Custom Product",
        "Google Shopping / Custom Label 0", "Variant Image", "Variant Weight Unit",
        "Cost per item", "Status",
    ]

    csv_path = os.path.join(out_dir, "shopify_import_casalivo_es.csv")
    json_path = os.path.join(out_dir, "shopify_import_casalivo_es.json")
    qa_path = os.path.join(out_dir, "import_qa_report.json")

    products, payloads, skipped, warnings_out = [], [], [], []

    for inv in inventory:
        sku = (inv.get("Seller SKU") or "").strip()
        if not sku:
            continue
        qty = int(num(inv.get("Qty Available")))
        list_price = money(inv.get("List Price (Excl. Tax)") or inv.get("List Price"))
        shipping_fee = money(inv.get("ES Shipping Fee") or inv.get("ES Shipping Fee (Excl. Tax)") or inv.get("Shipping Fee") or 0)
        cost = list_price + shipping_fee
        if qty <= SAFETY_BUFFER:
            skipped.append({"sku": sku, "reason": f"qty {qty} <= buffer {SAFETY_BUFFER}"})
            continue
        if cost <= 0:
            skipped.append({"sku": sku, "reason": "no price"})
            continue
        hit = catalog.get(sku)
        if not hit:
            skipped.append({"sku": sku, "reason": "not in catalog"})
            continue
        row, idx = hit["row"], hit["idx"]

        raw_name = str(row[idx.get("Product Name (Spanish)")] or row[idx.get("Product Name")] or "").strip()
        title = clean_title(raw_name)
        if not title:
            skipped.append({"sku": sku, "reason": "empty title"})
            continue
        images = []
        for n in range(1, 13):
            col = f"Product Image {n}"
            if col in idx:
                url = str(row[idx[col]] or "").strip()
                if url.startswith("http") and url not in images:
                    images.append(url)
        if not images:
            skipped.append({"sku": sku, "reason": "no image"})
            continue
        images = images[:MAX_IMAGES]

        brand_sku = f"{SKU_PREFIX}{sku.upper()}"
        price = pricing(cost)
        avail = max(0, qty - SAFETY_BUFFER)
        grams = int(num(row[idx.get("Product Weight (kg)", -1)]) * 1000)
        category_en = str(row[idx.get("Category", -1)] or "").strip()
        category = category_en or "Hogar"
        category_es = spanish_group(category_en)
        body = build_body(row, idx, sku)
        handle = f"{slug(brand_sku)}-{slug(raw_name)}"
        category_std = shopify_standard_category(category_en)

        # Requirement 1: English 3rd-level category name into tags
        tags = [t for t in [BRAND, "Envío desde Europa", "Stock verificado", category_en, category_es] if t]
        seo_title = title[:70]
        seo_desc = scrub(raw_name)[:160]

        main = {
            "Handle": handle,
            "Title": title,
            "Body (HTML)": body,
            "Vendor": BRAND,
            "Product Category": category_std,
            "Type": category,
            "Tags": ", ".join(dict.fromkeys(tags)),
            "Published": "TRUE",
            "Option1 Name": "Title",
            "Option1 Value": "Default Title",
            "Variant SKU": brand_sku,
            "Variant Grams": str(grams),
            "Variant Inventory Tracker": "shopify",
            "Variant Inventory Qty": str(avail),
            "Variant Inventory Policy": "deny",
            "Variant Fulfillment Service": "manual",
            "Variant Price": f"{price['price']:.2f}",
            "Variant Compare At Price": "",  # GMC Protection: 默认不填划线价防虚假判定
            "Variant Requires Shipping": "TRUE",
            "Variant Taxable": "TRUE",
            "Variant Barcode": generate_ean13_barcode(sku),
            "Image Src": images[0],
            "Image Position": "1",
            "Image Alt Text": title,
            "Gift Card": "FALSE",
            "SEO Title": seo_title,
            "SEO Description": seo_desc,
            "Google Shopping / MPN": brand_sku,
            "Google Shopping / Condition": "new",
            "Google Shopping / Custom Product": "FALSE",
            "Google Shopping / Custom Label 0": BRAND,
            "Variant Image": images[0],
            "Variant Weight Unit": "kg",
            "Cost per item": f"{cost:.2f}",
            "Status": "active",
        }
        products.append(main)
        for n, url in enumerate(images[1:], start=2):
            products.append({"Handle": handle, "Image Src": url, "Image Position": str(n),
                             "Image Alt Text": f"{title} - vista {n}"})

        payloads.append({
            "sku": brand_sku,
            "source_sku": sku,
            "title": title,
            "body_html": body,
            "vendor": BRAND,
            "product_type": category,
            "tags": tags,
            "status": "active",
            "stock": avail,
            "cost_eur": round(cost, 2),
            "images": [{"src": u} for u in images],
            "variants": [{
                "sku": brand_sku,
                "price": f"{price['price']:.2f}",
                "compare_at_price": f"{price['compare_at_price']:.2f}",
                "inventory_management": "shopify",
                "inventory_policy": "deny",
                "weight": round(grams / 1000, 2),
                "weight_unit": "kg",
                "requires_shipping": True,
                "taxable": True,
            }],
            "metafields": [{"namespace": "casalivo", "key": "cost_eur", "type": "number_decimal",
                            "value": f"{cost:.2f}"}],
        })

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=fieldnames)
        w.writeheader()
        for p in products:
            w.writerow({k: p.get(k, "") for k in fieldnames})

    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(payloads, fh, ensure_ascii=False, indent=2)

    # ---------------- QA
    leaks = []
    for m in products:
        if not m.get("Title"):
            continue
        blob = " ".join(str(m.get(k, "")) for k in ("Title", "Body (HTML)", "Tags", "SEO Title", "SEO Description"))
        low = blob.lower()
        for bad in BANNED + ["aosom", "homcom", "outsunny", "pawhut", "vinsetto"]:
            if bad.lower() in low:
                leaks.append({"sku": m.get("Variant SKU"), "leak": bad})
                break
    costs = [money(p["Cost per item"]) for p in products if p.get("Variant SKU")]
    prices = [float(p["Variant Price"]) for p in products if p.get("Variant SKU")]
    report = {
        "inventory_rows": len(inventory),
        "published_products": len(payloads),
        "csv_rows": len(products),
        "skipped": skipped,
        "whitewash_leaks": leaks,
        "price_min": min(prices) if prices else 0,
        "price_max": max(prices) if prices else 0,
        "cost_total": round(sum(costs), 2),
        "price_total": round(sum(prices), 2),
        "avg_margin_pct": round(sum((p - c) / p * 100 for c, p in zip(costs, prices)) / len(prices), 1) if prices else 0,
        "sources": {"inventory": inv_path, "catalog": cat_path},
        "warnings": warnings_out,
    }
    with open(qa_path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)

    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"\nCSV  -> {csv_path}\nJSON -> {json_path}")


if __name__ == "__main__":
    main()
